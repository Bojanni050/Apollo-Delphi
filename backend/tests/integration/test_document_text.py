"""The text of a document for the reading pane: the extracted text, its pages and where its fragments lie."""
import io
from pathlib import Path

import pytest

from app.core.config import get_settings
from app.services.documents import reading
from app.services.extraction.base import ExtractedPage, ExtractionResult

@pytest.fixture(autouse=True)
def _fresh_cache():
    reading._cache.clear()  # document ids restart in every test; a remembered text must not leak across them
    yield
    reading._cache.clear()


MD = "# Harbour\nThe harbour budget is 250000 EUR.\nQuay repairs are included.\n\n# Sponsor\nThe project sponsor is Maria Chen.\n"


def _upload(client, name="notes.md", text=MD):
    ws = client.post("/api/workspaces", json={"name": "w"}).json()["id"]
    doc = client.post(
        "/api/documents", params={"workspace_id": ws}, files={"file": (name, io.BytesIO(text.encode()), "text/plain")}
    ).json()
    return ws, doc


def test_the_text_comes_back_with_line_count_and_no_pages_for_a_text_file(client):
    _, doc = _upload(client)
    body = client.get(f"/api/documents/{doc['id']}/text").json()
    assert body["text"] == MD.strip() and body["line_count"] == MD.strip().count("\n") + 1
    assert body["filename"] == "notes.md" and body["file_type"] == "md" and body["pages"] == [] and body["truncated"] is False
    assert body["chunks"] == [], "nothing is indexed yet"


def test_indexed_fragments_are_listed_with_lines_that_hold_their_text(client):
    _, doc = _upload(client)
    client.post(f"/api/documents/{doc['id']}/index")
    body = client.get(f"/api/documents/{doc['id']}/text").json()
    lines = body["text"].split("\n")
    assert [c["section"] for c in body["chunks"]] == ["Harbour", "Sponsor"]
    for chunk in body["chunks"]:
        covered = "\n".join(lines[chunk["line_start"] - 1 : chunk["line_end"]])
        assert chunk["section"] in covered
    assert (body["chunks"][0]["line_start"], body["chunks"][1]["line_start"]) == (1, 5)


def test_unknown_documents_and_vanished_files_are_404(client):
    assert client.get("/api/documents/99999/text").status_code == 404
    _, doc = _upload(client)
    for path in Path(get_settings().upload_dir).glob("*"):
        path.unlink()
    resp = client.get(f"/api/documents/{doc['id']}/text")
    assert resp.status_code == 404 and "gone" in resp.json()["detail"]


def test_an_unreadable_file_is_a_422_with_the_reason(client):
    ws = client.post("/api/workspaces", json={"name": "w"}).json()["id"]
    doc = client.post(
        "/api/documents", params={"workspace_id": ws}, files={"file": ("broken.pdf", io.BytesIO(b"%PDF-1.4 nope"), "application/pdf")}
    ).json()
    resp = client.get(f"/api/documents/{doc['id']}/text")
    assert resp.status_code == 422 and resp.json()["detail"]


def test_page_starts_are_the_lines_where_each_page_begins():
    extraction = ExtractionResult(
        pages=[ExtractedPage(1, "alpha\nbeta"), ExtractedPage(2, "gamma"), ExtractedPage(3, "delta\nepsilon\nzeta")],
        full_text="alpha\nbeta\n\ngamma\n\ndelta\nepsilon\nzeta",
    )
    readable = reading.readable_text(extraction)
    assert [(p.page_number, p.line) for p in readable.pages] == [(1, 1), (2, 4), (3, 6)]
    assert readable.line_count == 8
    lines = readable.text.split("\n")
    assert lines[3] == "gamma" and lines[5] == "delta"


def test_a_page_less_text_has_no_page_markers():
    assert reading.readable_text(ExtractionResult(pages=[ExtractedPage(None, "x")], full_text="x")).pages == []


def test_a_very_long_text_is_cut_and_flagged(monkeypatch):
    monkeypatch.setattr(reading, "MAX_CHARS", 10)
    readable = reading.readable_text(ExtractionResult(pages=[ExtractedPage(None, "a" * 50)], full_text="a" * 50))
    assert readable.truncated is True and len(readable.text) == 10


def test_extraction_is_remembered_per_document_and_content(client, monkeypatch):
    _, doc = _upload(client)
    from app.services.documents.service import document_service

    calls = []
    real = document_service.extract
    monkeypatch.setattr(document_service, "extract", lambda d: (calls.append(d.id), real(d))[1])
    reading._cache.clear()
    for _ in range(3):
        assert client.get(f"/api/documents/{doc['id']}/text").status_code == 200
    assert len(calls) == 1, "the second and third reading come from memory"


def _docx_bytes():
    from docx import Document as Word

    word = Word()
    word.add_heading("Harbour plan", level=1)
    para = word.add_paragraph("The budget is ")
    para.add_run("250000 EUR").bold = True
    word.add_paragraph("Quay repairs", style="List Bullet")
    table = word.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Cost"
    buf = io.BytesIO()
    word.save(buf)
    return buf.getvalue()


def _upload_bytes(client, name, data):
    ws = client.post("/api/workspaces", json={"name": "w"}).json()["id"]
    return client.post("/api/documents", params={"workspace_id": ws}, files={"file": (name, io.BytesIO(data), "application/octet-stream")}).json()


def test_the_original_file_is_served_inline_with_its_own_type(client):
    _, doc = _upload(client)
    resp = client.get(f"/api/documents/{doc['id']}/file")
    assert resp.status_code == 200 and resp.content == MD.encode()
    assert resp.headers["content-type"].startswith("text/markdown")
    assert resp.headers["content-disposition"].startswith("inline") and resp.headers["x-content-type-options"] == "nosniff"
    assert client.get("/api/documents/99999/file").status_code == 404


def test_a_pdf_is_served_as_application_pdf(client):
    doc = _upload_bytes(client, "paper.pdf", b"%PDF-1.4 minimal")
    resp = client.get(f"/api/documents/{doc['id']}/file")
    assert resp.status_code == 200 and resp.headers["content-type"] == "application/pdf" and resp.content.startswith(b"%PDF")


def test_a_word_document_is_converted_to_html_with_its_structure(client):
    doc = _upload_bytes(client, "plan.docx", _docx_bytes())
    body = client.get(f"/api/documents/{doc['id']}/html").json()
    html = body["html"]
    assert "<h1>Harbour plan</h1>" in html and "<strong>250000 EUR</strong>" in html
    assert "<li>Quay repairs</li>" in html and "<table>" in html


def test_only_word_documents_are_converted_to_html(client):
    _, doc = _upload(client)
    assert client.get(f"/api/documents/{doc['id']}/html").status_code == 415
